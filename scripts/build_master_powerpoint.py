# -*- coding: utf-8 -*-
"""
Build 100% Pixel-Perfect Master PowerPoint from the 14 High-Res 1080p Rendered Slides
"""
import os
from pptx import Presentation
from pptx.util import Inches

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SLIDES_DIR = os.path.join(BASE_DIR, "presentation", "slides_rendered")
OUTPUT_PPTX = os.path.join(BASE_DIR, "presentation", "Lusail_University_Master_Pitch.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

for i in range(1, 15):
    img_path = os.path.join(SLIDES_DIR, f"slide_{i}.png")
    if os.path.exists(img_path):
        slide = prs.slides.add_slide(blank_layout)
        # Add full-bleed image
        slide.shapes.add_picture(img_path, 0, 0, Inches(13.333), Inches(7.5))
        print(f"Added Slide {i} to Master PowerPoint")

prs.save(OUTPUT_PPTX)
print(f"Master PowerPoint saved successfully at: {OUTPUT_PPTX}")
